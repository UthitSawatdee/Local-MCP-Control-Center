from __future__ import annotations

import csv
from pathlib import Path

from .conftest import add_scope, allow_capabilities, enable_tools


def _apply_document_edit(broker, tool: str, args: dict) -> dict:
    result = broker.invoke(tool, args)
    assert result["status"] == "ok", result
    assert result["executed"] is True
    return result


def test_csv_transform_is_allowlisted_and_neutralizes_formula(broker, workspace: Path) -> None:
    target = workspace / "items.csv"
    with target.open("w", encoding="utf-8", newline="") as handle:
        csv.writer(handle).writerows([["name", "value"], ["one", "1"]])
    add_scope(broker, workspace)
    allow_capabilities(broker, "test-scope", "read", "write")
    enable_tools(broker, "csv_transform")

    _apply_document_edit(
        broker,
        "csv_transform",
        {
            "scope_id": "test-scope",
            "relative_path": "items.csv",
            "operation": "append_rows",
            "rows": [["two", "=NETWORK()"]],
        },
    )

    with target.open(encoding="utf-8", newline="") as handle:
        rows = list(csv.reader(handle))
    assert rows[-1] == ["two", "'=NETWORK()"]


def test_read_csv_returns_bounded_rows_with_encoding_and_secret_redaction(broker, workspace: Path) -> None:
    target = workspace / "items.csv"
    with target.open("w", encoding="utf-8-sig", newline="") as handle:
        csv.writer(handle, delimiter=";").writerows(
            [
                ["name", "value", "notes"],
                ["one", "1", "ok"],
                ["two", "password=top-secret", "ok"],
                ["three", "3", "ok"],
            ]
        )
    before = target.read_bytes()
    add_scope(broker, workspace)
    allow_capabilities(broker, "test-scope", "read")

    result = broker.invoke(
        "read_csv",
        {
            "scope_id": "test-scope",
            "relative_path": "items.csv",
            "start_row": 1,
            "max_rows": 3,
            "max_columns": 2,
            "delimiter": ";",
        },
    )

    assert result["status"] == "ok", result
    assert result["format"] == "csv"
    assert result["encoding"] == "utf-8-sig"
    assert result["delimiter"] == ";"
    assert result["row_count"] == 4
    assert result["column_count"] == 3
    assert result["start_row"] == 1
    assert result["end_row"] == 3
    assert result["rows"] == [
        ["name", "value"],
        ["one", "1"],
        ["two", "[REDACTED]"],
    ]
    assert result["rows_truncated"] is True
    assert result["columns_truncated"] is True
    assert result["truncated"] is True
    assert result["has_more"] is True
    assert result["content_hash"].startswith("sha256:")
    assert target.read_bytes() == before
    assert broker.audit_page(1)[0]["tool"] == "read_csv"


def test_read_csv_supports_row_pagination_and_is_allowed_in_read_batch(broker, workspace: Path) -> None:
    target = workspace / "items.csv"
    with target.open("w", encoding="utf-8", newline="") as handle:
        csv.writer(handle).writerows([["id"], ["one"], ["two"]])
    add_scope(broker, workspace)
    allow_capabilities(broker, "test-scope", "read")

    page = broker.invoke(
        "read_csv",
        {"scope_id": "test-scope", "relative_path": "items.csv", "start_row": 2, "max_rows": 2},
    )
    batch = broker.invoke(
        "tool_batch",
        {
            "scope_id": "test-scope",
            "operations": [
                {"operation": "read_csv", "arguments": {"relative_path": "items.csv", "max_rows": 1}}
            ],
        },
    )

    assert page["status"] == "ok", page
    assert page["rows"] == [["one"], ["two"]]
    assert page["start_row"] == 2
    assert page["end_row"] == 3
    assert page["has_more"] is False
    assert batch["status"] == "ok", batch
    assert batch["succeeded"] == 1
    assert batch["results"][0]["result"]["rows"] == [["id"]]


def test_xlsx_and_docx_operations_use_format_adapters(broker, workspace: Path) -> None:
    from docx import Document
    from openpyxl import Workbook, load_workbook

    workbook_path = workspace / "book.xlsx"
    workbook = Workbook()
    workbook.active.title = "Data"
    workbook.save(workbook_path)

    document_path = workspace / "letter.docx"
    document = Document()
    document.add_paragraph("Hello customer")
    document.save(document_path)

    add_scope(broker, workspace)
    allow_capabilities(broker, "test-scope", "read", "write")
    enable_tools(broker, "xlsx_edit", "docx_edit")

    _apply_document_edit(
        broker,
        "xlsx_edit",
        {
            "scope_id": "test-scope",
            "relative_path": "book.xlsx",
            "operation": "update_cells",
            "updates": [{"sheet": "Data", "cell": "A1", "value": "Approved"}],
        },
    )
    saved = load_workbook(workbook_path, data_only=False)
    assert saved["Data"]["A1"].value == "Approved"

    _apply_document_edit(
        broker,
        "docx_edit",
        {
            "scope_id": "test-scope",
            "relative_path": "letter.docx",
            "operation": "replace_text",
            "old": "customer",
            "new": "partner",
        },
    )
    edited = Document(document_path)
    assert edited.paragraphs[0].text == "Hello partner"
