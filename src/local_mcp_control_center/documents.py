from __future__ import annotations

import csv
import io
import math
import re
from pathlib import Path
from typing import Any

from .errors import PolicyError
from .filesystem import redact_text


CELL_RE = re.compile(r"^[A-Z]{1,3}[1-9][0-9]*$")
INVALID_SHEET_NAME_CHARS = set("[]:*?/\\")
CSV_DELIMITERS = (",", ";", "\t", "|")
MAX_CSV_BYTES = 20 * 1024 * 1024
MAX_CSV_START_ROW = 50_000_000
MAX_CSV_ROWS = 1_000
MAX_CSV_COLUMNS = 200


def _safe_cell(cell: str) -> str:
    if not isinstance(cell, str) or not CELL_RE.fullmatch(cell.upper()):
        raise PolicyError("INVALID_INPUT", f"invalid spreadsheet cell: {cell!r}")
    return cell.upper()


def _safe_value(value: Any, *, allow_formulas: bool = False) -> Any:
    if isinstance(value, (str, int, float, bool)) or value is None:
        if isinstance(value, float) and not math.isfinite(value):
            raise PolicyError("INVALID_INPUT", "document numeric values must be finite")
        if isinstance(value, str) and value.startswith(("=", "+", "-", "@")) and not allow_formulas:
            return "'" + value
        return value
    raise PolicyError("INVALID_INPUT", "document values must be scalar")


def _safe_sheet_name(value: Any) -> str:
    if not isinstance(value, str):
        raise PolicyError("INVALID_INPUT", "worksheet name must be text")
    name = value.strip()
    if not 1 <= len(name) <= 31 or any(char in INVALID_SHEET_NAME_CHARS for char in name):
        raise PolicyError("INVALID_INPUT", "worksheet name must be 1-31 characters without []:*?/\\")
    return name


class DocumentAdapter:
    """Format-specific adapter with explicit operations and no script evaluation."""

    def read_csv(
        self,
        target: Path,
        *,
        start_row: int = 1,
        max_rows: int = 200,
        max_columns: int = 50,
        delimiter: str = ",",
        max_bytes: int = MAX_CSV_BYTES,
    ) -> dict[str, Any]:
        """Read a bounded, redacted slice of a UTF-8 CSV file."""

        if not isinstance(start_row, int) or isinstance(start_row, bool) or not 1 <= start_row <= MAX_CSV_START_ROW:
            raise PolicyError("INVALID_INPUT", f"start_row must be between 1 and {MAX_CSV_START_ROW}")
        if not isinstance(max_rows, int) or isinstance(max_rows, bool) or not 1 <= max_rows <= MAX_CSV_ROWS:
            raise PolicyError("INVALID_INPUT", f"max_rows must be between 1 and {MAX_CSV_ROWS}")
        if not isinstance(max_columns, int) or isinstance(max_columns, bool) or not 1 <= max_columns <= MAX_CSV_COLUMNS:
            raise PolicyError("INVALID_INPUT", f"max_columns must be between 1 and {MAX_CSV_COLUMNS}")
        if delimiter not in CSV_DELIMITERS:
            raise PolicyError("INVALID_INPUT", "delimiter must be one of ',', ';', tab, or '|'")

        rows, encoding, newline = self._read_csv(
            target,
            max_bytes=max_bytes,
            delimiter=delimiter,
        )
        total_rows = len(rows)
        total_columns = max((len(row) for row in rows), default=0)
        start_index = start_row - 1
        selected_rows = rows[start_index:start_index + max_rows]
        returned_rows = [
            [redact_text(value) for value in row[:max_columns]]
            for row in selected_rows
        ]
        rows_truncated = start_index + len(selected_rows) < total_rows
        columns_truncated = total_columns > max_columns
        return {
            "format": "csv",
            "encoding": encoding,
            "newline": newline,
            "delimiter": delimiter,
            "row_count": total_rows,
            "column_count": total_columns,
            "start_row": start_row,
            "end_row": start_row + len(selected_rows) - 1 if selected_rows else None,
            "rows": returned_rows,
            "truncated": rows_truncated or columns_truncated,
            "rows_truncated": rows_truncated,
            "columns_truncated": columns_truncated,
            "has_more": rows_truncated,
        }

    def inspect_csv(self, target: Path) -> dict[str, Any]:
        rows, encoding, newline = self._read_csv(target)
        return {
            "format": "csv",
            "rows": len(rows),
            "columns": max((len(row) for row in rows), default=0),
            "encoding": encoding,
            "newline": repr(newline),
            "preview": [row[:20] for row in rows[:10]],
        }

    def transform_csv(self, target: Path, payload: dict[str, Any]) -> tuple[bytes, dict[str, Any]]:
        rows, encoding, newline = self._read_csv(target)
        operation = payload.get("operation")
        allow_formulas = bool(payload.get("allow_formulas", False))
        if operation == "append_rows":
            additions = payload.get("rows")
            self._validate_rows(additions)
            rows.extend([[_safe_value(value, allow_formulas=allow_formulas) for value in row] for row in additions])
        elif operation == "update_cells":
            updates = payload.get("updates")
            if not isinstance(updates, list) or len(updates) > 1000:
                raise PolicyError("INVALID_INPUT", "updates must be a list of at most 1000 cells")
            for item in updates:
                if not isinstance(item, dict):
                    raise PolicyError("INVALID_INPUT", "each CSV update must be an object")
                row_number = item.get("row")
                column_number = item.get("column")
                if not isinstance(row_number, int) or not isinstance(column_number, int) or row_number < 1 or column_number < 1:
                    raise PolicyError("INVALID_INPUT", "CSV row and column are one-based positive integers")
                while len(rows) < row_number:
                    rows.append([])
                row = rows[row_number - 1]
                while len(row) < column_number:
                    row.append("")
                row[column_number - 1] = _safe_value(item.get("value"), allow_formulas=allow_formulas)
        elif operation == "sort_rows":
            column_number = payload.get("column")
            if not isinstance(column_number, int) or column_number < 1:
                raise PolicyError("INVALID_INPUT", "CSV sort column must be a positive integer")
            header = bool(payload.get("header", True))
            start = 1 if header and rows else 0
            rows[start:] = sorted(
                rows[start:],
                key=lambda row: str(row[column_number - 1] if len(row) >= column_number else "").casefold(),
                reverse=bool(payload.get("descending", False)),
            )
        elif operation == "rename_columns":
            mapping = payload.get("mapping")
            if not isinstance(mapping, dict) or not rows:
                raise PolicyError("INVALID_INPUT", "mapping and a header row are required")
            rows[0] = [
                _safe_value(mapping.get(value, value), allow_formulas=allow_formulas)
                for value in rows[0]
            ]
        else:
            raise PolicyError("INVALID_INPUT", "unsupported CSV operation")

        output = io.StringIO(newline="")
        writer = csv.writer(output, lineterminator=newline)
        writer.writerows(rows)
        text = output.getvalue()
        raw = text.encode("utf-8-sig" if encoding == "utf-8-sig" else "utf-8")
        return raw, {"format": "csv", "operation": operation, "rows": len(rows), "columns": max((len(row) for row in rows), default=0)}

    @staticmethod
    def _validate_rows(rows: Any) -> None:
        if not isinstance(rows, list) or len(rows) > 10_000:
            raise PolicyError("INVALID_INPUT", "rows must be a list of at most 10000 rows")
        for row in rows:
            if not isinstance(row, list) or len(row) > 500:
                raise PolicyError("INVALID_INPUT", "each row must contain at most 500 values")

    @staticmethod
    def _read_csv(
        target: Path,
        *,
        max_bytes: int = MAX_CSV_BYTES,
        delimiter: str = ",",
    ) -> tuple[list[list[str]], str, str]:
        if not isinstance(max_bytes, int) or isinstance(max_bytes, bool) or max_bytes < 1:
            raise PolicyError("INVALID_INPUT", "CSV byte limit must be a positive integer")
        if delimiter not in CSV_DELIMITERS:
            raise PolicyError("INVALID_INPUT", "delimiter must be one of ',', ';', tab, or '|'")
        limit = min(max_bytes, MAX_CSV_BYTES)
        if target.stat().st_size > limit:
            raise PolicyError("QUOTA_EXCEEDED", f"CSV exceeds the {limit} byte read limit")
        raw = target.read_bytes()
        if len(raw) > limit:
            raise PolicyError("QUOTA_EXCEEDED", f"CSV exceeds the {limit} byte read limit")
        encoding = "utf-8-sig" if raw.startswith(b"\xef\xbb\xbf") else "utf-8"
        try:
            text = raw.decode(encoding)
        except UnicodeDecodeError as exc:
            raise PolicyError("FORMAT_UNSUPPORTED", "CSV must be UTF-8 or UTF-8 with BOM") from exc
        newline = "\r\n" if "\r\n" in text else "\n"
        reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
        try:
            rows = [row for row in reader]
        except csv.Error as exc:
            raise PolicyError("FORMAT_VALIDATION_FAILED", f"unable to parse CSV: {exc}") from exc
        return rows, encoding, newline

    def edit_xlsx(self, target: Path, payload: dict[str, Any]) -> tuple[bytes, dict[str, Any]]:
        if target.suffix.lower() not in {".xlsx", ".xlsm"}:
            raise PolicyError("FORMAT_UNSUPPORTED", "only .xlsx and .xlsm are writable")
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise PolicyError("FORMAT_UNSUPPORTED", "openpyxl is not installed") from exc

        keep_vba = target.suffix.lower() == ".xlsm"
        try:
            workbook = load_workbook(target, keep_vba=keep_vba, data_only=False)
        except Exception as exc:
            raise PolicyError("FORMAT_VALIDATION_FAILED", f"unable to open workbook: {exc}") from exc
        operation = payload.get("operation")
        changed = 0
        if operation == "update_cells":
            updates = payload.get("updates")
            if not isinstance(updates, list) or len(updates) > 1000:
                raise PolicyError("INVALID_INPUT", "updates must be a list of at most 1000 cells")
            for item in updates:
                if not isinstance(item, dict) or not isinstance(item.get("sheet"), str):
                    raise PolicyError("INVALID_INPUT", "XLSX update requires sheet, cell, and value")
                sheet = item["sheet"]
                if sheet not in workbook.sheetnames:
                    raise PolicyError("TARGET_NOT_FOUND", f"worksheet not found: {sheet}")
                cell = _safe_cell(item.get("cell", ""))
                workbook[sheet][cell] = _safe_value(item.get("value"), allow_formulas=bool(payload.get("allow_formulas", False)))
                changed += 1
        elif operation == "append_rows":
            sheet = payload.get("sheet")
            rows = payload.get("rows")
            if not isinstance(sheet, str) or sheet not in workbook.sheetnames:
                raise PolicyError("TARGET_NOT_FOUND", "worksheet not found")
            self._validate_rows(rows)
            for row in rows:
                workbook[sheet].append([_safe_value(value, allow_formulas=bool(payload.get("allow_formulas", False))) for value in row])
                changed += 1
        elif operation == "rename_sheet":
            old_name, new_name = payload.get("old_name"), payload.get("new_name")
            if not isinstance(old_name, str) or old_name not in workbook.sheetnames:
                raise PolicyError("TARGET_NOT_FOUND", "worksheet to rename was not found")
            new_name = _safe_sheet_name(new_name)
            if new_name in workbook.sheetnames:
                raise PolicyError("INVALID_INPUT", "worksheet name already exists")
            workbook[old_name].title = new_name
            changed = 1
        elif operation == "add_sheet":
            name = payload.get("name")
            name = _safe_sheet_name(name)
            if name in workbook.sheetnames:
                raise PolicyError("INVALID_INPUT", "worksheet name already exists")
            workbook.create_sheet(name)
            changed = 1
        else:
            raise PolicyError("INVALID_INPUT", "unsupported XLSX operation")

        buffer = io.BytesIO()
        try:
            workbook.save(buffer)
        except Exception as exc:
            raise PolicyError("FORMAT_VALIDATION_FAILED", f"unable to save workbook: {exc}") from exc
        return buffer.getvalue(), {"format": target.suffix.lower().lstrip("."), "operation": operation, "changed": changed, "sheet_count": len(workbook.sheetnames)}

    def edit_docx(self, target: Path, payload: dict[str, Any]) -> tuple[bytes, dict[str, Any]]:
        if target.suffix.lower() != ".docx":
            raise PolicyError("FORMAT_UNSUPPORTED", "only .docx is writable")
        try:
            from docx import Document
        except ImportError as exc:
            raise PolicyError("FORMAT_UNSUPPORTED", "python-docx is not installed") from exc
        try:
            document = Document(target)
        except Exception as exc:
            raise PolicyError("FORMAT_VALIDATION_FAILED", f"unable to open document: {exc}") from exc

        operation = payload.get("operation")
        changed = 0
        if operation == "replace_text":
            old, new = payload.get("old"), payload.get("new")
            if not isinstance(old, str) or not old or not isinstance(new, str) or len(new) > 50_000:
                raise PolicyError("INVALID_INPUT", "old and new text are required")
            for paragraph in document.paragraphs:
                if old in paragraph.text:
                    paragraph.text = paragraph.text.replace(old, new)
                    changed += 1
            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if old in cell.text:
                            cell.text = cell.text.replace(old, new)
                            changed += 1
        elif operation == "append_paragraph":
            text = payload.get("text")
            if not isinstance(text, str) or len(text) > 50_000:
                raise PolicyError("INVALID_INPUT", "paragraph text is invalid or too large")
            document.add_paragraph(text)
            changed = 1
        elif operation == "update_table_cell":
            table_index, row_index, column_index, text = (
                payload.get("table"), payload.get("row"), payload.get("column"), payload.get("text")
            )
            if not all(isinstance(value, int) and value >= 0 for value in (table_index, row_index, column_index)) or not isinstance(text, str) or len(text) > 50_000:
                raise PolicyError("INVALID_INPUT", "table, row, column, and text are required")
            try:
                document.tables[table_index].rows[row_index].cells[column_index].text = text
            except IndexError as exc:
                raise PolicyError("TARGET_NOT_FOUND", "document table cell not found") from exc
            changed = 1
        else:
            raise PolicyError("INVALID_INPUT", "unsupported DOCX operation")

        buffer = io.BytesIO()
        try:
            document.save(buffer)
        except Exception as exc:
            raise PolicyError("FORMAT_VALIDATION_FAILED", f"unable to save document: {exc}") from exc
        return buffer.getvalue(), {"format": "docx", "operation": operation, "changed": changed}
